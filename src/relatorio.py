"""
relatorio.py — Gera o Relatório Executivo (DOCX e PDF) a partir do modelo do Google Docs.

Como funciona:
  1. Abre a cópia do modelo (docs/modelo_relatorio/modelo_google_docs.docx), que já
     tem capa, folha de rosto, sumário e os títulos das seções.
  2. Lê os NÚMEROS gerados pelos notebooks (reports/tables/*.csv e numeros_chave.json)
     e as FIGURAS (reports/figures/*.png). Nenhum resultado é digitado à mão.
  3. Escreve o texto de cada seção logo abaixo do título correspondente.
  4. Converte para PDF com o LibreOffice e atualiza o sumário com as páginas reais.

Uso:  python -m src.relatorio
"""

import copy
import json
import re
import shutil
import subprocess
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src import config
from src.inferencia import formatar_p

SAIDA_DOCX = config.PASTA_REPORTS / "relatorio_executivo.docx"
SAIDA_PDF = config.PASTA_REPORTS / "relatorio_executivo.pdf"
URL_REPOSITORIO = config.URL_REPOSITORIO
FONTE = "Fonte: elaboração própria a partir de IBM HR Analytics (Kaggle, 2017)."


# ---------------------------------------------------------------------------
# Formatação de números no padrão brasileiro
# ---------------------------------------------------------------------------
def br(valor: float, casas: int = 1) -> str:
    """1234.5 → '1.234,5' (milhar com ponto, decimal com vírgula)."""
    texto = f"{valor:,.{casas}f}"
    return texto.replace(",", "§").replace(".", ",").replace("§", ".")


def pct(valor: float, casas: int = 1) -> str:
    return f"{br(valor, casas)}%"


def p_br(p: float) -> str:
    return formatar_p(p).replace("< 0,001", "< 0,001")


# ---------------------------------------------------------------------------
# Motor de escrita: insere elementos em sequência a partir de um parágrafo-âncora
# ---------------------------------------------------------------------------
class Escritor:
    """
    O python-docx só sabe adicionar conteúdo no FIM do documento. Esta classe cria
    cada elemento no fim e o move para logo depois do "cursor", permitindo
    escrever dentro das seções do modelo, na ordem certa.
    """

    def __init__(self, doc: Document):
        self.doc = doc
        self.cursor = None
        self.n_fig = 0
        self.n_tab = 0

    def posicionar(self, paragrafo) -> None:
        self.cursor = paragrafo._p

    def _mover(self, elemento) -> None:
        self.cursor.addnext(elemento)
        self.cursor = elemento

    # --- texto -----------------------------------------------------------
    def _runs(self, par, texto: str, tamanho: float = 12, cor=None) -> None:
        """Escreve texto com **negrito** e *itálico* no estilo Markdown."""
        for parte in re.split(r"(\*\*.+?\*\*|\*.+?\*)", texto):
            if not parte:
                continue
            if parte.startswith("**"):
                run = par.add_run(parte[2:-2]); run.bold = True
            elif parte.startswith("*"):
                run = par.add_run(parte[1:-1]); run.italic = True
            else:
                run = par.add_run(parte)
            run.font.size = Pt(tamanho)
            run.font.name = "Arial"
            if cor:
                run.font.color.rgb = RGBColor.from_string(cor)

    def paragrafo(self, texto: str, recuo: bool = True, tamanho: float = 12,
                  alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, espaco: float = 1.5, depois: float = 6):
        par = self.doc.add_paragraph()
        self._runs(par, texto, tamanho)
        fmt = par.paragraph_format
        fmt.alignment = alinhamento
        fmt.line_spacing = espaco
        fmt.space_after = Pt(depois)
        fmt.space_before = Pt(0)
        fmt.first_line_indent = Cm(1.25) if recuo else None
        self._mover(par._p)
        return par

    def topicos(self, itens: list, tamanho: float = 12) -> None:
        for item in itens:
            par = self.doc.add_paragraph()
            self._runs(par, "•  " + item, tamanho)
            fmt = par.paragraph_format
            fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt.line_spacing = 1.5
            fmt.left_indent = Cm(1.0)
            fmt.first_line_indent = Cm(-0.5)
            fmt.space_after = Pt(3)
            self._mover(par._p)

    def titulo2(self, texto: str) -> None:
        par = self.doc.add_paragraph(style="Heading 2")
        run = par.add_run(texto)
        run.bold = True
        run.font.size = Pt(12)
        par.paragraph_format.space_before = Pt(12)
        par.paragraph_format.space_after = Pt(6)
        self._mover(par._p)

    def legenda(self, texto: str, acima: bool = True) -> None:
        """Legenda no padrão ABNT: título acima (Figura/Tabela N – ...) e fonte abaixo."""
        par = self.doc.add_paragraph()
        self._runs(par, texto, tamanho=10)
        fmt = par.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER if acima else WD_ALIGN_PARAGRAPH.LEFT
        fmt.line_spacing = 1.0
        fmt.space_before = Pt(6 if acima else 2)
        fmt.space_after = Pt(2 if acima else 10)
        fmt.keep_with_next = acima
        self._mover(par._p)

    # --- figuras e tabelas -----------------------------------------------
    def figura(self, arquivo: str, titulo: str, largura: float = 16.0, nota: str | None = None) -> int:
        self.n_fig += 1
        self.legenda(f"**Figura {self.n_fig}** – {titulo}")
        par = self.doc.add_paragraph()
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_after = Pt(0)
        par.add_run().add_picture(str(config.PASTA_FIGURAS / f"{arquivo}.png"), width=Cm(largura))
        self._mover(par._p)
        self.legenda((nota + " " if nota else "") + FONTE, acima=False)
        return self.n_fig

    def tabela(self, df: pd.DataFrame, titulo: str, larguras: list | None = None,
               tamanho: float = 9, nota: str | None = None) -> int:
        """Tabela com cabeçalho destacado e apenas bordas horizontais (estilo ABNT/IBGE)."""
        self.n_tab += 1
        self.legenda(f"**Tabela {self.n_tab}** – {titulo}")
        tab = self.doc.add_table(rows=1, cols=len(df.columns))
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordas_horizontais(tab)
        for j, col in enumerate(df.columns):
            cel = tab.rows[0].cells[j]
            cel.text = ""
            self._runs(cel.paragraphs[0], f"**{col}**", tamanho)
            _sombrear(cel, "E8EEF7")
        for _, linha in df.iterrows():
            cels = tab.add_row().cells
            for j, valor in enumerate(linha):
                cels[j].text = ""
                self._runs(cels[j].paragraphs[0], str(valor), tamanho)
        for linha in tab.rows:
            for j, cel in enumerate(linha.cells):
                for par in cel.paragraphs:
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.space_before = Pt(1)
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if larguras:
                    cel.width = Cm(larguras[j])
        if larguras:
            # o LibreOffice usa as larguras da grade da tabela (tblGrid), não só as das células
            tab.autofit = False
            for j, col in enumerate(tab.columns):
                col.width = Cm(larguras[j])
        _repetir_cabecalho(tab.rows[0])
        for linha in tab.rows:  # impede que uma linha seja cortada entre duas páginas
            linha._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        self._mover(tab._tbl)
        self.legenda((nota + " " if nota else "") + FONTE, acima=False)
        return self.n_tab

    def quebra_pagina(self) -> None:
        par = self.doc.add_paragraph()
        par.add_run().add_break(WD_BREAK.PAGE)
        self._mover(par._p)


def _bordas_horizontais(tabela) -> None:
    tbl_pr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for nome in ("top", "bottom", "insideH"):
        b = OxmlElement(f"w:{nome}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6" if nome == "insideH" else "10")
        b.set(qn("w:color"), "BFBFBF" if nome == "insideH" else "404040")
        bordas.append(b)
    tbl_pr.append(bordas)


def _sombrear(celula, cor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(shd)


def _repetir_cabecalho(linha) -> None:
    tr_pr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _corrigir_margens(doc: Document) -> None:
    """O Google Docs exporta margens com casas decimais (ex.: 1700.78), que o Word rejeita."""
    for pg_mar in doc.element.body.iter(qn("w:pgMar")):
        for atributo, valor in list(pg_mar.attrib.items()):
            try:
                pg_mar.set(atributo, str(round(float(valor))))
            except ValueError:
                pass


def _ajustar_capa(doc: Document) -> None:
    """
    No modelo, a capa e a folha de rosto ultrapassam uma página (o próprio Google Docs
    joga "Penedo - AL / 2026" para a página seguinte). Zerar o espaço após as linhas
    em branco faz cada uma caber em uma página, sem mudar o texto.
    """
    for par in doc.paragraphs:
        if par.text.strip() == "SUMÁRIO":
            break
        tem_quebra = any(br.get(qn("w:type")) == "page" for br in par._p.iter(qn("w:br")))
        tem_imagem = any(True for _ in par._p.iter(qn("w:drawing")))
        if not par.text.strip() and not tem_quebra and not tem_imagem:
            par.paragraph_format.space_after = Pt(0)
    # remove até 3 linhas em branco antes de "Penedo - AL" (capa e folha de rosto)
    pars = doc.paragraphs
    for i, par in enumerate(pars):
        if par.text.strip() == "SUMÁRIO":
            break
        if par.text.strip() == "Penedo - AL":
            for anterior in pars[max(0, i - 3):i]:
                if not anterior.text.strip():
                    anterior._p.getparent().remove(anterior._p)


def _justificar_folha_de_rosto(doc: Document) -> None:
    """
    Na folha de rosto, o texto da natureza do trabalho ("Trabalho solicitado...") e o
    nome do docente ficam JUSTIFICADOS. O recuo esquerdo do modelo (≈ 8 cm) é mantido:
    só o alinhamento muda.
    """
    inicios = ("Trabalho solicitado na disciplina", "Orientado pelo docente")
    for par in doc.paragraphs:
        if par.text.strip().startswith(inicios):
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _achar(doc: Document, texto: str, estilo: str | None = None):
    for par in doc.paragraphs:
        if par.text.strip() == texto and (estilo is None or par.style.name == estilo):
            return par
    raise ValueError(f"Parágrafo '{texto}' não encontrado no modelo")


# ---------------------------------------------------------------------------
# Carregamento dos resultados
# ---------------------------------------------------------------------------
def carregar_resultados() -> dict:
    t = config.PASTA_TABELAS
    return {
        "num": json.loads((t / "numeros_chave.json").read_text(encoding="utf-8")),
        "decisoes": pd.read_csv(t / "decisoes_limpeza.csv"),
        "tipagem": pd.read_csv(t / "tipagem_variaveis.csv"),
        "descritiva": pd.read_csv(t / "descritiva_numericas.csv", index_col=0),
        "testes": pd.read_csv(t / "resultados_inferencia.csv"),
        "ics": pd.read_csv(t / "intervalos_confianca.csv"),
        "renda_nivel": pd.read_csv(t / "renda_mediana_por_nivel_status.csv"),
        "freq": pd.read_csv(t / "frequencias_categoricas.csv"),
    }


def _traduzir_variaveis(texto: str) -> str:
    for en, pt in sorted(config.NOMES_PT.items(), key=lambda kv: -len(kv[0])):
        texto = re.sub(rf"\b{re.escape(en)}\b", pt.split(" (")[0], texto)
    return texto.replace("Rotatividade", "Saída")


# ---------------------------------------------------------------------------
# Conteúdo do relatório
# ---------------------------------------------------------------------------
def escrever_resumo(doc, e: Escritor, R: dict) -> None:
    n = R["num"]
    e.posicionar(_achar(doc, "RESUMO"))
    e.paragrafo(
        f"Este relatório investiga os fatores associados à rotatividade voluntária de funcionários a partir da base "
        f"IBM HR Analytics ({br(n['n'], 0)} funcionários, 35 variáveis). O objetivo é apoiar a área de Gestão de Pessoas "
        f"na priorização de ações de retenção com base em evidências estatísticas. Após o diagnóstico de qualidade e a "
        f"limpeza dos dados, foram aplicadas estatística descritiva e cinco procedimentos inferenciais (intervalo de "
        f"confiança, teste t, ANOVA, qui-quadrado e teste de correlação, com α = 0,05). A taxa de rotatividade é de "
        f"**{pct(n['taxa_geral'])}** (IC 95%: {pct(n['taxa_ic'][0])} a {pct(n['taxa_ic'][1])}). O **principal achado** é "
        f"que funcionários que fazem hora extra saem **{br(n['he_rr'], 1)} vezes mais** ({pct(n['taxa_he_sim'])} contra "
        f"{pct(n['taxa_he_nao'])}; p < 0,001). Baixo envolvimento com o trabalho ({pct(n['inv_baixo'])} de saída), baixa "
        f"satisfação com o ambiente, equilíbrio vida-trabalho ruim e cargos de entrada, de menor renda, também estão "
        f"associados à saída. Já o reajuste salarial depende apenas da nota de desempenho, que não diferencia quem sai "
        f"de quem fica. A **principal recomendação** é implantar a gestão ativa de horas extras, combinada a um programa "
        f"de retenção para cargos de entrada e a um sistema de alerta baseado em pesquisas de clima.",
        recuo=False, espaco=1.0, depois=8)
    # KPIs em destaque
    kpis = pd.DataFrame([
        [pct(n["taxa_geral"]), f"{br(n['he_rr'], 1)}×", pct(n["inv_baixo"]), pct(n["taxa_q1_renda"])],
        ["taxa de rotatividade", "risco de saída com hora extra", "saída com envolvimento baixo",
         "saída no quartil de menor renda"],
    ], columns=["Rotatividade", "Hora extra", "Envolvimento", "Renda"])
    e.tabela(kpis, "Indicadores-chave do estudo", tamanho=10, larguras=[4, 4, 4, 4])
    kw = _achar(doc, "Palavras-chave:")
    e._runs(kw, "Rotatividade de pessoal; People Analytics; Teste de hipóteses; Inteligência empresarial; Gestão de pessoas.")

    e.posicionar(_achar(doc, "ABSTRACT"))
    e.paragrafo(
        f"This report investigates the factors associated with voluntary employee attrition using the IBM HR Analytics "
        f"dataset ({br(n['n'], 0).replace('.', ',')} employees, 35 variables), aiming to support People Management in "
        f"prioritizing evidence-based retention actions. After data quality assessment and cleaning, descriptive "
        f"statistics and five inferential procedures were applied (confidence interval, t-test, ANOVA, chi-square and "
        f"correlation test, α = 0.05). The attrition rate is {n['taxa_geral']:.1f}% (95% CI: {n['taxa_ic'][0]:.1f}% to "
        f"{n['taxa_ic'][1]:.1f}%). The main finding is that employees who work overtime leave {n['he_rr']:.1f} times "
        f"more often ({n['taxa_he_sim']:.1f}% vs. {n['taxa_he_nao']:.1f}%; p < 0.001). Low job involvement, low "
        f"environment satisfaction, poor work-life balance and lower-paid entry-level roles are also associated with "
        f"attrition, while salary raises depend only on the performance rating, which does not distinguish leavers "
        f"from stayers. The main recommendation is active overtime management, combined with a retention program for "
        f"entry-level roles and an early-warning system based on engagement surveys.",
        recuo=False, espaco=1.0, depois=8)
    kw = _achar(doc, "Keywords:")
    e._runs(kw, "Employee attrition; People Analytics; Hypothesis testing; Business intelligence; Human resources.")


def escrever_introducao(doc, e: Escritor, R: dict) -> None:
    n = R["num"]
    e.posicionar(_achar(doc, "1. Introdução e Contexto"))
    e.titulo2("1.1 Motivação")
    e.paragrafo(
        "A rotatividade voluntária de funcionários (*attrition*) é um dos indicadores mais acompanhados pela área de "
        "Gestão de Pessoas. Cada desligamento gera custos diretos (recrutamento, seleção, integração e treinamento do "
        "substituto) e indiretos (perda de conhecimento, queda de produtividade e sobrecarga da equipe). Entender **quais "
        "fatores estão associados à saída** permite que a empresa atue de forma preventiva e direcione recursos para os "
        "grupos de maior risco, em vez de aplicar ações genéricas a todo o quadro.")
    e.paragrafo(
        "Este trabalho aplica o ciclo de Inteligência Empresarial (coleta, tratamento, análise descritiva, inferência "
        "estatística e recomendação) a uma base de Recursos Humanos, com foco em três dimensões controláveis pela "
        "gestão: **remuneração**, **envolvimento com o trabalho** e **ambiente de trabalho**.")
    e.titulo2("1.2 Base de dados")
    e.paragrafo(
        f"Foi utilizada a base **IBM HR Analytics Employee Attrition & Performance**, publicada no Kaggle "
        f"({config.URL_DATASET}). Ela contém **{br(n['n'], 0)} registros** (um por funcionário) e **35 variáveis**: "
        f"demográficas (idade, gênero, estado civil), de carreira (cargo, nível, tempo de empresa, promoções), de "
        f"remuneração (renda mensal, reajuste, ações) e de percepção (satisfação com o ambiente, com o trabalho, "
        f"envolvimento e equilíbrio vida-trabalho, em escalas de 1 a 4). A variável-alvo é **Attrition** "
        f"(saiu: sim/não), com {br(n['saidas'], 0)} saídas registradas.")
    e.paragrafo(
        "A base atende aos requisitos da atividade (mínimo de 500 registros, 5 variáveis numéricas e 3 categóricas). "
        "**Importante:** os dados são **fictícios**, criados por cientistas de dados da IBM para fins didáticos. As "
        "conclusões ilustram o método analítico e não descrevem uma empresa real.")
    e.titulo2("1.3 Objetivo geral e perguntas de negócio")
    e.paragrafo(
        "**Objetivo geral:** identificar os fatores salariais, de envolvimento e de ambiente de trabalho associados à "
        "rotatividade de funcionários, para subsidiar ações de retenção priorizadas por evidência estatística. Para "
        "isso, foram formuladas cinco perguntas de negócio (Tabela 2).")
    perguntas = pd.DataFrame([
        ["Q1", "Qual é a taxa de rotatividade e ela difere entre departamentos?", "IC 95% da proporção; Qui-Quadrado"],
        ["Q2", "Funcionários que saem ganham menos do que os que ficam?", "Teste t; IC 95% da diferença de médias"],
        ["Q3", "Envolvimento, ambiente, satisfação e equilíbrio vida-trabalho estão associados à saída?",
         "Qui-Quadrado (×4) com correção de Holm"],
        ["Q4", "Fazer hora extra está associado a mais saídas?", "Qui-Quadrado 2×2; IC 95% da diferença de proporções"],
        ["Q5", "O desempenho e o envolvimento se refletem no reajuste salarial?", "Teste t; ANOVA; Teste de correlação"],
    ], columns=["Nº", "Pergunta de negócio", "Procedimento inferencial"])
    e.tabela(perguntas, "Perguntas de negócio e procedimentos associados", larguras=[1.2, 8.8, 6.0])


def escrever_metodologia(doc, e: Escritor, R: dict) -> None:
    titulo = _achar(doc, "2. Metodologia")
    titulo.paragraph_format.page_break_before = True
    e.posicionar(titulo)
    e.titulo2("2.1 Ferramentas")
    e.paragrafo(
        "A análise foi desenvolvida em **Python 3.14**, com as bibliotecas **pandas** e **NumPy** (manipulação de dados), "
        "**SciPy** e **statsmodels** (testes estatísticos e intervalos de confiança), **Matplotlib** e **seaborn** "
        "(visualização) e **Jupyter Notebook** (documentação do passo a passo). O código está organizado em três "
        "notebooks (limpeza, descritiva e inferencial) e em módulos reutilizáveis na pasta *src/*. O script *run_all.py* "
        f"reproduz toda a análise e este relatório a partir do arquivo original. O repositório está em {URL_REPOSITORIO}.")
    e.titulo2("2.2 Etapas do tratamento dos dados")
    e.paragrafo("O fluxo de trabalho seguiu as etapas do roteiro da atividade (Figura 1).")
    e.figura("fig01_fluxo_metodologia", "Fluxo metodológico da análise")
    e.paragrafo(
        "O diagnóstico de qualidade não encontrou **valores ausentes** nem **registros duplicados**. Três colunas "
        "constantes foram removidas e três colunas de \"taxa\" foram excluídas da análise por inconsistência com a renda "
        "mensal (correlação próxima de zero). Os **outliers** foram identificados pela regra do intervalo interquartil "
        "(valores fora de Q1 − 1,5·IQR e Q3 + 1,5·IQR) e **mantidos**, por representarem situações legítimas, como "
        "salários de diretores e funcionários com muitos anos de casa (Figura 2). O efeito deles foi controlado na etapa "
        "inferencial. A Tabela 3 resume as decisões.")
    dec = R["decisoes"].rename(columns={"verificacao": "Verificação", "resultado": "Resultado",
                                        "acao": "Ação", "justificativa": "Justificativa"})
    e.tabela(dec, "Decisões de limpeza e qualidade dos dados", larguras=[3.0, 4.6, 3.4, 5.0], tamanho=8.5)
    e.figura("fig02_outliers_boxplots", "Boxplots das variáveis numéricas e outliers pela regra do IQR",
             nota="Pontos laranja: valores fora dos limites do IQR.")
    e.titulo2("2.3 Tipagem das variáveis")
    e.paragrafo(
        "Cada variável foi classificada pelo seu tipo estatístico, que define as medidas e testes adequados. As escalas "
        "de 1 a 4 (satisfação, envolvimento) foram tratadas como **qualitativas ordinais**, e não como medidas contínuas "
        "(Tabela 4).")
    tip = R["tipagem"].groupby("tipo_estatistico", sort=False)["descricao"].apply(
        lambda s: "; ".join(x.split(" (")[0] for x in s)).reset_index()
    tip.columns = ["Tipo estatístico", "Variáveis"]
    e.tabela(tip, "Tipagem estatística das variáveis analisadas", larguras=[3.6, 12.4], tamanho=8.5)
    e.titulo2("2.4 Procedimentos estatísticos")
    e.paragrafo(
        "Na etapa descritiva foram calculadas média, mediana, desvio padrão, intervalo interquartil (IQR), assimetria, "
        "mínimo e máximo das variáveis numéricas, e frequências absolutas e relativas das categóricas. A matriz de "
        "correlação usou o coeficiente de **Spearman**, adequado a variáveis assimétricas e ordinais.")
    e.paragrafo(
        "Na etapa inferencial, todos os testes adotaram **nível de significância α = 0,05** e seguiram o mesmo roteiro: "
        "hipóteses H₀ e H₁, verificação de pressupostos, escolha justificada do teste, estatística, p-valor, decisão, "
        "tamanho de efeito e interpretação de negócio. Os procedimentos foram:")
    e.topicos([
        "**Intervalos de confiança (95%)** para proporções pelo método de **Wilson** e para a diferença de proporções "
        "pelo método de **Newcombe**, ambos mais precisos que o método clássico para proporções distantes de 50% ou "
        "grupos pequenos;",
        "**Teste t de Student** para duas médias, na versão de **Welch** quando o teste de Levene indicou variâncias "
        "diferentes. A normalidade foi avaliada por Shapiro-Wilk e gráficos Q-Q. Com amostras grandes (n > 30), o "
        "Teorema Central do Limite sustenta a inferência sobre médias;",
        "**ANOVA** de um fator para três ou mais médias, com teste de Levene e post-hoc de Tukey quando necessário;",
        "**Teste Qui-Quadrado** de independência para pares de variáveis categóricas, verificando frequências "
        "esperadas ≥ 5, com **correção de Holm** para o conjunto de quatro testes da Q3;",
        "**Teste de correlação** (Pearson e Spearman) para relações entre variáveis quantitativas.",
    ])
    e.paragrafo(
        "Os tamanhos de efeito (V de Cramér, d de Cohen, η² e r) foram interpretados segundo Cohen (1988), para "
        "distinguir diferenças **estatisticamente significativas** de diferenças **relevantes na prática**.")


def escrever_descritiva(doc, e: Escritor, R: dict) -> None:
    n = R["num"]
    d = R["descritiva"]
    titulo = _achar(doc, "3. Análise Descritiva e Exploratória")
    titulo.paragraph_format.page_break_before = True
    e.posicionar(titulo)

    e.titulo2("3.1 Variáveis numéricas")
    tab = pd.DataFrame({
        "Variável": [config.NOMES_PT[v] for v in d.index],
        "Média": [br(x, 1) for x in d["media"]],
        "Mediana": [br(x, 1) for x in d["mediana"]],
        "DP": [br(x, 1) for x in d["desvio_padrao"]],
        "IQR": [br(x, 1) for x in d["IQR"]],
        "Assim.": [br(x, 2) for x in d["assimetria"]],
        "Mín.": [br(x, 0) for x in d["minimo"]],
        "Máx.": [br(x, 0) for x in d["maximo"]],
    })
    r = d.loc["MonthlyIncome"]
    e.paragrafo(
        f"A Tabela 5 apresenta as medidas-resumo. A **idade** é praticamente simétrica (média de {br(d.loc['Age', 'media'])} "
        f"e mediana de {br(d.loc['Age', 'mediana'], 0)} anos). A **renda mensal** tem forte assimetria à direita "
        f"(assimetria = {br(r['assimetria'], 2)}): a média (US$ {br(r['media'], 0)}) supera bastante a mediana "
        f"(US$ {br(r['mediana'], 0)}), por isso a mediana representa melhor o salário típico. As variáveis de tempo "
        f"(anos na empresa, desde a última promoção) também são assimétricas: a maioria dos funcionários tem pouco tempo "
        f"de casa, e poucos têm décadas (Figura 3).")
    e.tabela(tab, "Medidas descritivas das variáveis numéricas (n = 1.470)",
             larguras=[4.6, 1.6, 1.6, 1.6, 1.5, 1.5, 1.4, 1.6], tamanho=8.5,
             nota="DP = desvio padrão; IQR = intervalo interquartil; Assim. = coeficiente de assimetria.")
    e.figura("fig03_histogramas", "Distribuição das variáveis numéricas",
             nota="Linha laranja: média; linha preta tracejada: mediana.")

    e.titulo2("3.2 Variáveis categóricas")
    f = R["freq"]
    def fr(var, cat):
        lin = f[(f["variavel"] == var) & (f["categoria"].astype(str) == str(cat))].iloc[0]
        return pct(lin["freq_relativa_pct"])
    e.paragrafo(
        f"A variável-alvo é desbalanceada: **{pct(n['taxa_geral'])} dos funcionários saíram**. P&D concentra "
        f"{fr('Department', 'Research & Development')} do quadro, Vendas {fr('Department', 'Sales')} e RH apenas "
        f"{fr('Department', 'Human Resources')}. Fazem hora extra {fr('OverTime', 'Yes')} dos funcionários. Um ponto de "
        f"atenção é a **avaliação de desempenho**: só existem as notas 3 ({fr('PerformanceRating', 3)}) e 4 "
        f"({fr('PerformanceRating', 4)}), e nenhuma nota 1 ou 2, o que sugere **inflação de avaliações** (Figura 4).")
    e.figura("fig04_frequencias_categoricas", "Frequência relativa das principais variáveis categóricas")

    e.titulo2("3.3 Rotatividade por departamento e cargo (Q1)")
    e.paragrafo(
        f"Vendas tem a maior taxa de saída entre os departamentos, mas a diferença é muito maior entre **cargos**: "
        f"**Representantes de Vendas perdem {pct(n['taxa_rep_vendas'])}** do quadro e Técnicos de Laboratório "
        f"{pct(n['taxa_tec_lab'])}, enquanto cargos de gestão ficam abaixo de 7% (Figura 5). O problema está concentrado "
        f"em **cargos de entrada**, e não em departamentos inteiros.")
    e.figura("fig05_q1_departamento_cargo", "Taxa de saída por departamento e por cargo",
             nota="Barras de erro: IC 95% (Wilson). Linha tracejada: taxa geral.")

    e.titulo2("3.4 Renda e rotatividade (Q2)")
    rn = R["renda_nivel"].set_index("JobLevel")
    e.paragrafo(
        f"Quem saiu tem renda mediana de **US$ {br(n['renda_mediana_saiu'], 0)}**, contra **US$ "
        f"{br(n['renda_mediana_ficou'], 0)}** de quem ficou. No quartil de menor renda a taxa de saída é "
        f"**{pct(n['taxa_q1_renda'])}**, cerca de três vezes a do quartil mais alto ({pct(n['taxa_q4_renda'])}) "
        f"(Figura 6). Porém, **dentro do mesmo nível de cargo, a renda de quem sai e de quem fica é quase igual** "
        f"(nível 1: US$ {br(rn.loc[1, 'Saiu'], 0)} × US$ {br(rn.loc[1, 'Ficou'], 0)}; nível 2: US$ "
        f"{br(rn.loc[2, 'Saiu'], 0)} × US$ {br(rn.loc[2, 'Ficou'], 0)}), e {br(rn.loc[1, 'n_Saiu'] / n['saidas'] * 100, 0)}% "
        f"das saídas estão no nível 1. A renda baixa funciona, em boa parte, como um **marcador de início de carreira**.")
    e.figura("fig06_q2_renda", "Renda mensal e rotatividade",
             nota="◆ = média. Barras de erro: IC 95% (Wilson).")

    e.titulo2("3.5 Envolvimento e ambiente de trabalho (Q3)")
    e.paragrafo(
        f"Nas quatro escalas de percepção, **o nível mais baixo concentra o risco** (Figura 7). O envolvimento tem o "
        f"gradiente mais claro: a taxa cai de **{pct(n['inv_baixo'])}** (baixo) para **{pct(n['inv_muito_alto'])}** "
        f"(muito alto). Satisfação baixa com o ambiente leva a {pct(n['amb_baixo'])} de saída, e equilíbrio vida-trabalho "
        f"\"ruim\" a {pct(n['wlb_ruim'])}. Nos demais níveis, as taxas ficam próximas da média.")
    e.figura("fig07_q3_envolvimento_ambiente", "Taxa de saída por nível de envolvimento, satisfação e equilíbrio vida-trabalho",
             nota="Barras de erro: IC 95% (Wilson). Laranja: maior taxa.")

    e.titulo2("3.6 Hora extra (Q4)")
    e.paragrafo(
        f"Hora extra é a variável com a maior diferença entre grupos: **{pct(n['taxa_he_sim'])}** de saída entre quem faz "
        f"hora extra, contra **{pct(n['taxa_he_nao'])}** entre quem não faz. O efeito aparece em todos os quartis de "
        f"renda e é extremo no quartil mais baixo, onde mais da metade de quem faz hora extra saiu (Figura 8). "
        f"**Hora extra e salário baixo somam risco.**")
    e.figura("fig08_q4_hora_extra", "Hora extra e rotatividade, no total e por quartil de renda")

    e.titulo2("3.7 Desempenho, envolvimento e reajuste (Q5)")
    e.paragrafo(
        f"O reajuste salarial segue uma **regra fixa atrelada à nota**: toda nota 3 recebeu entre 11% e 19%, e toda nota 4 "
        f"entre 20% e 25%. O reajuste médio **não varia** com o nível de envolvimento nem com a renda (Figura 9). Além "
        f"disso, a taxa de saída é praticamente igual nas duas notas ({pct(n['taxa_nota3'])} × {pct(n['taxa_nota4'])}): "
        f"**a avaliação de desempenho não diferencia quem sai de quem fica.**")
    e.figura("fig09_q5_desempenho_reajuste", "Avaliação de desempenho, envolvimento e reajuste salarial")

    e.titulo2("3.8 Correlações e perfil de risco")
    e.paragrafo(
        "Na matriz de correlação (Figura 10), a saída tem correlação positiva mais forte com **hora extra** (ρ ≈ 0,25) e "
        "negativas com **nível do cargo, renda, experiência, tempo de casa e idade** (ρ entre −0,17 e −0,20). Todas são "
        "fracas individualmente, o que indica que **a rotatividade é multifatorial**. Renda, nível do cargo e "
        "experiência são fortemente correlacionados entre si (ρ de 0,71 a 0,92), ou seja, medem a senioridade. As "
        "escalas de satisfação **não se correlacionam entre si**: são dimensões independentes, e cada uma exige ação "
        "própria.")
    e.figura("fig10_matriz_correlacao", "Matriz de correlação de Spearman", largura=15.5)
    e.paragrafo(
        "O perfil complementar (Figura 11) confirma que o risco se concentra no **início da jornada**: jovens de 18 a 25 "
        "anos, funcionários com até 2 anos de casa e solteiros têm as maiores taxas de saída. Quem viaja com frequência "
        "também sai mais.")
    e.figura("fig11_perfil_complementar", "Taxa de saída por perfil demográfico e de carreira")


def _linha_teste(testes: pd.DataFrame, chave: str) -> pd.Series:
    return testes[testes["variaveis"].str.contains(chave, regex=False)].iloc[0]


def escrever_inferencial(doc, e: Escritor, R: dict) -> None:
    n = R["num"]
    T = R["testes"]
    titulo = _achar(doc, "4. Análise Inferencial")
    titulo.paragraph_format.page_break_before = True
    e.posicionar(titulo)
    e.paragrafo(
        "Esta seção verifica se as diferenças observadas na análise descritiva são **estatisticamente significativas** ou "
        "podem ser atribuídas ao acaso. Em todos os testes, α = 0,05: se p-valor < 0,05, rejeita-se H₀. Por se tratar de "
        "estudo observacional, os resultados indicam **associação, e não causalidade**.")

    e.titulo2("4.1 Q1: Taxa de rotatividade e departamento")
    dep = _linha_teste(T, "Department ×")
    cargo = _linha_teste(T, "JobRole ×")
    e.paragrafo(
        f"**Intervalo de confiança.** A taxa de rotatividade estimada é de {pct(n['taxa_geral'])}, com **IC 95% (Wilson) "
        f"de {pct(n['taxa_ic'][0])} a {pct(n['taxa_ic'][1])}**: de cada 100 funcionários, entre 14 e 18 deixam a empresa. "
        f"Esse valor é a linha de base para metas de retenção.")
    e.paragrafo(
        f"**Qui-Quadrado.** H₀: departamento e rotatividade são independentes; H₁: estão associados. O pressuposto de "
        f"frequências esperadas ≥ 5 foi atendido (menor = 10,2). Resultado: χ²(2) = {br(n['dep_chi2'], 2)}, "
        f"p = {p_br(n['dep_p'])}. **Rejeita-se H₀**: a taxa difere entre departamentos. Porém o V de Cramér = "
        f"{br(n['dep_v'], 3)} é desprezível. O teste complementar por **cargo** também rejeita H₀ (χ² = "
        f"{br(n['cargo_chi2'], 1)}, p < 0,001), com efeito quase três vezes maior (V = {br(n['cargo_v'], 3)}). "
        f"**Interpretação:** a retenção deve ser direcionada por **cargo**, e não por departamento.")

    e.titulo2("4.2 Q2: Renda de quem sai × quem fica")
    q2 = _linha_teste(T, "MonthlyIncome por grupo")
    q2log = _linha_teste(T, "log(MonthlyIncome)")
    e.paragrafo(
        "H₀: μ(renda | saiu) = μ(renda | ficou); H₁: as médias são diferentes (bilateral). **Pressupostos:** o "
        "Shapiro-Wilk rejeitou a normalidade nos dois grupos (renda assimétrica, ver Apêndice A), mas com n = 237 e "
        "n = 1.233 o Teorema Central do Limite garante a normalidade aproximada da média amostral. O teste de Levene "
        "rejeitou a igualdade de variâncias (p < 0,001), então foi aplicado o **teste t de Welch**.")
    e.paragrafo(
        f"**Resultado:** t({br(n['renda_gl'], 1)}) = {br(n['renda_t'], 2)}, p < 0,001. **Rejeita-se H₀.** Quem sai ganha "
        f"em média **US$ {br(abs(n['renda_diff']), 0)} a menos por mês** (IC 95%: US$ {br(abs(n['renda_diff_ic'][1]), 0)} "
        f"a US$ {br(abs(n['renda_diff_ic'][0]), 0)} a menos), com efeito pequeno a médio (d de Cohen = "
        f"{br(n['renda_d'], 2)}). Como verificação de robustez, o teste t sobre o **logaritmo da renda** (escala em que "
        f"as variâncias são iguais e a distribuição é quase simétrica) leva à mesma conclusão ({q2log['estatistica'].replace('.', ',')}, "
        f"p < 0,001). **Interpretação:** a diferença existe, mas reflete principalmente a concentração das saídas em "
        f"cargos de entrada (Seção 3.4). O foco deve ser a atratividade da carreira inicial.")

    e.titulo2("4.3 Q3: Envolvimento, ambiente, satisfação e vida-trabalho")
    rot = {"JobInvolvement × Attrition": "Envolvimento com o trabalho",
           "EnvironmentSatisfaction × Attrition": "Satisfação com o ambiente",
           "JobSatisfaction × Attrition": "Satisfação com o trabalho",
           "WorkLifeBalance × Attrition": "Equilíbrio vida-trabalho"}
    tab = pd.DataFrame([[nome, f"{br(n['q3_chi2'][k], 2)}", "3", p_br(n["q3_p_holm"][k]),
                         br(n["q3_v"][k], 3), "Rejeita H₀"] for k, nome in rot.items()],
                       columns=["Fator", "χ²", "gl", "p (Holm)", "V de Cramér", "Decisão"])
    e.paragrafo(
        "Para cada fator: H₀: o fator e a rotatividade são independentes; H₁: estão associados. Foram aplicados quatro "
        "testes Qui-Quadrado (todas as frequências esperadas ≥ 12,9). Como quatro testes simultâneos aumentam a chance "
        "de falso positivo, os p-valores foram ajustados pela **correção de Holm** (Tabela 6).")
    e.tabela(tab, "Testes Qui-Quadrado dos fatores de envolvimento e ambiente × rotatividade",
             larguras=[5.0, 1.8, 1.0, 2.4, 2.4, 2.6])
    e.paragrafo(
        f"**Os quatro testes rejeitam H₀**, mesmo após o ajuste. O **envolvimento** tem o maior efeito "
        f"(V = {br(n['q3_v']['JobInvolvement × Attrition'], 3)}): {pct(n['inv_baixo'])} de saída no nível baixo, contra "
        f"{pct(n['inv_muito_alto'])} no muito alto. Os efeitos individuais são pequenos, mas, como as quatro dimensões "
        f"são independentes entre si, **o risco se acumula** em quem tem várias dimensões no nível baixo. "
        f"**Interpretação:** pesquisas de clima que identifiquem respostas \"baixo\" funcionam como alerta precoce de saída.")

    e.titulo2("4.4 Q4: Hora extra")
    reducao = n["n_he_sim"] * n["he_diff"] / 100
    e.paragrafo(
        f"H₀: hora extra e rotatividade são independentes; H₁: estão associadas. Foi aplicado o **Qui-Quadrado 2×2** "
        f"(com correção de Yates), complementado pelo **IC 95% da diferença de proporções** (Newcombe). Resultado: "
        f"χ²(1) = {br(n['he_chi2'], 2)}, p < 0,001. **Rejeita-se H₀.** A taxa de saída é {pct(n['taxa_he_sim'])} com hora "
        f"extra e {pct(n['taxa_he_nao'])} sem hora extra: diferença de **{br(n['he_diff'], 1)} pontos percentuais** "
        f"(IC 95%: {br(n['he_diff_ic'][0], 1)} a {br(n['he_diff_ic'][1], 1)} p.p.) e **risco relativo de "
        f"{br(n['he_rr'], 2)}**. É a associação mais forte do estudo (V = {br(n['he_v'], 3)}). **Interpretação:** se a "
        f"taxa de quem faz hora extra caísse ao nível de quem não faz, seriam cerca de **{br(reducao, 0)} saídas a menos** "
        f"({br(reducao / n['saidas'] * 100, 0)}% do total). É um potencial, já que a relação é de associação.")

    e.titulo2("4.5 Q5: Desempenho, envolvimento e reajuste")
    e.paragrafo(
        f"**(a) Teste t (nota 4 × nota 3).** H₀: o reajuste médio é igual nas duas notas. Levene rejeitou a igualdade de "
        f"variâncias, então foi usado o **t de Welch** (n = 226 e 1.244; TCL aplicável). Resultado: t = "
        f"{br(n['reajuste_t'], 1)}, p < 0,001. **Rejeita-se H₀**: nota 4 recebe em média {br(n['reajuste_nota4'], 1)}% "
        f"contra {br(n['reajuste_nota3'], 1)}% (diferença com IC 95% de {br(n['reajuste_diff_ic'][0], 1)} a "
        f"{br(n['reajuste_diff_ic'][1], 1)} p.p.). As faixas não se sobrepõem: há uma regra fixa.")
    e.paragrafo(
        f"**(b) ANOVA (reajuste × 4 níveis de envolvimento).** H₀: as médias são iguais nos quatro níveis. Levene não "
        f"rejeitou a igualdade de variâncias (p = 0,72), então foi usada a ANOVA clássica (todos os grupos com n ≥ 83). "
        f"Resultado: F(3; 1466) = {br(n['anova_f'], 2)}, p = {p_br(n['anova_p'])}. **Não se rejeita H₀** "
        f"(η² = {br(n['anova_eta2'], 3)}): o envolvimento **não é recompensado** no reajuste.")
    e.paragrafo(
        f"**(c) Correlação (renda × reajuste).** H₀: ρ = 0. Como a renda é assimétrica e com outliers, o coeficiente "
        f"principal é o de **Spearman**: ρ = {br(n['corr_rho'], 3)}, p = {p_br(n['corr_p'])}. **Não se rejeita H₀.** "
        f"**Interpretação da Q5:** o reajuste depende apenas da nota, e a nota não está associada à saída. A principal "
        f"ferramenta de recompensa **não atua sobre os fatores que explicam a rotatividade**.")

    e.titulo2("4.6 Síntese dos resultados inferenciais")
    tab = T.copy()
    tab["Pergunta"] = tab["pergunta"].str.split(" · ").str[0] + tab["pergunta"].str.contains("Robustez").map(
        {True: " (robustez)", False: ""})
    tab["Teste"] = tab["teste"].str.replace(" de independência", "").str.replace(" (correção de Yates)", " (Yates)")
    tab["Variáveis"] = tab["variaveis"].map(_traduzir_variaveis).str.replace(" por grupo", "")
    tab["Estatística"] = tab["estatistica"].str.replace(".", ",", regex=False)
    tab["p-valor"] = [p_br(p if pd.isna(ph) else ph) + (" (Holm)" if not pd.isna(ph) else "")
                      for p, ph in zip(tab["p_valor"], tab["p_valor_holm"])]
    tab["Decisão"] = tab["decisao"].str.replace("H0", "H₀")
    tab["Efeito"] = (tab["efeito"].str.replace(r" \(variância compartilhada\)", "", regex=True)
                     .str.replace(r" \(eta quadrado\)", "", regex=True).str.replace(".", ",", regex=False))
    e.tabela(tab[["Pergunta", "Teste", "Variáveis", "Estatística", "p-valor", "Decisão", "Efeito"]],
             "Resumo de todos os testes de hipóteses (α = 0,05)",
             larguras=[1.9, 2.5, 3.5, 2.0, 1.9, 1.8, 2.6], tamanho=7.5)
    ics = R["ics"].copy()
    ics["Estimativa"] = [f"{br(v, 1)} {u}" for v, u in zip(ics["estimativa"], ics["unidade"])]
    ics["IC 95%"] = [f"[{br(a, 1)}; {br(b, 1)}]" for a, b in zip(ics["ic_inferior"], ics["ic_superior"])]
    ics = ics.rename(columns={"pergunta": "Pergunta", "parametro": "Parâmetro", "metodo": "Método"})
    e.tabela(ics[["Pergunta", "Parâmetro", "Estimativa", "IC 95%", "Método"]],
             "Intervalos de confiança de 95%", larguras=[2.0, 6.8, 2.4, 2.8, 2.0], tamanho=8.5)
    e.paragrafo("A Figura 12 compara visualmente os intervalos das taxas de saída dos principais grupos.")
    e.figura("fig13_intervalos_confianca", "Taxa de saída estimada e IC 95% por grupo", largura=14.5,
             nota="Intervalos que não se cruzam indicam grupos com taxas claramente diferentes.")


def escrever_recomendacoes(doc, e: Escritor, R: dict) -> None:
    n = R["num"]
    titulo = _achar(doc, "5. Recomendações Práticas")
    titulo.paragraph_format.page_break_before = True
    e.posicionar(titulo)
    e.titulo2("5.1 Plano de ação")
    e.paragrafo(
        "As recomendações abaixo estão ordenadas por prioridade, considerando o tamanho do efeito encontrado e o grau de "
        "controle da gestão sobre cada fator. **Cada recomendação está ligada à evidência que a sustenta** (Tabela 9).")
    reducao = n["n_he_sim"] * n["he_diff"] / 100
    recs = pd.DataFrame([
        ["1. Gestão ativa de horas extras: limite mensal por pessoa, redistribuição de carga e revisão de quadro nas "
         "equipes com hora extra recorrente",
         f"Q4: χ² = {br(n['he_chi2'], 1)}, p < 0,001; RR = {br(n['he_rr'], 2)}; IC da diferença "
         f"{br(n['he_diff_ic'][0], 1)}–{br(n['he_diff_ic'][1], 1)} p.p. (Fig. 8, Seção 4.4)",
         "% do quadro com hora extra; taxa de saída de quem faz hora extra", "Alta"],
        ["2. Programa de retenção para cargos de entrada (Representante de Vendas, Técnico de Laboratório, nível 1): "
         "integração estruturada, mentoria, trilha de progressão e revisão do salário de entrada",
         f"Q1: cargo V = {br(n['cargo_v'], 2)} (vs. {br(n['dep_v'], 2)} do departamento); Q2: t de Welch, p < 0,001; "
         f"renda igual dentro do nível (Fig. 5, 6 e 11)",
         "Taxa de saída nos 2 primeiros anos; taxa por cargo", "Alta"],
        ["3. Sistema de alerta precoce: pesquisa de clima periódica e ação individual com quem responde \"baixo\" em "
         "envolvimento, ambiente, satisfação ou vida-trabalho",
         "Q3: 4 testes χ² com Holm, todos p < 0,002; envolvimento baixo = "
         f"{pct(n['inv_baixo'])} (Fig. 7, Tabela 6)",
         "% de respostas \"baixo\"; taxa de saída desse grupo", "Média-alta"],
        ["4. Revisão da política de reconhecimento: incluir envolvimento e engajamento nos critérios de reajuste e "
         "calibrar a avaliação de desempenho (hoje só notas 3 e 4)",
         f"Q5: t (nota → reajuste) p < 0,001; ANOVA p = {p_br(n['anova_p'])}; correlação p = {p_br(n['corr_p'])}; "
         f"saída igual nas duas notas (Fig. 9)",
         "Distribuição das notas; reajuste médio por nível de envolvimento", "Média"],
    ], columns=["Recomendação", "Evidência", "Indicador de acompanhamento", "Prioridade"])
    e.tabela(recs, "Plano de ação: recomendação, evidência, indicador e prioridade",
             larguras=[5.3, 5.0, 3.6, 2.1], tamanho=8)
    e.paragrafo(
        f"**Recomendação 1 (prioridade alta).** Hora extra é o fator de maior efeito e o que está **sob controle mais "
        f"direto** da gestão. Se a taxa de saída de quem faz hora extra caísse ao nível de quem não faz, o potencial "
        f"seria de cerca de **{br(reducao, 0)} saídas a menos por ciclo**, o que levaria a taxa geral de "
        f"{pct(n['taxa_geral'])} para cerca de {pct((n['saidas'] - reducao) / n['n'] * 100)}.")
    e.paragrafo(
        "**Recomendação 2 (prioridade alta).** As saídas se concentram em cargos de entrada, em funcionários jovens e com "
        "pouco tempo de casa. Como, no mesmo nível, quem sai não ganha menos que os pares, o problema é a **atratividade "
        "da carreira inicial** como um todo (salário de entrada, perspectiva de crescimento e integração), e não reajustes "
        "individuais.")
    e.paragrafo(
        "**Recomendação 3 (prioridade média-alta).** Envolvimento, ambiente, satisfação e equilíbrio vida-trabalho são "
        "dimensões **independentes**, e o risco se concentra no nível mais baixo de cada uma. Um painel que cruze as "
        "respostas da pesquisa de clima permite **agir antes do pedido de demissão**.")
    e.paragrafo(
        "**Recomendação 4 (prioridade média).** Hoje o reajuste depende só da nota de desempenho, e a nota não "
        "diferencia quem sai. Incluir o envolvimento nos critérios de reconhecimento alinha a política de remuneração aos "
        "fatores que de fato estão associados à permanência.")
    e.titulo2("5.2 Limitações e próximos passos")
    e.topicos([
        "Os dados são **fictícios** (simulados pela IBM): as recomendações demonstram o método e devem ser validadas com "
        "dados reais da organização;",
        "O estudo é **observacional** e de corte transversal: os testes indicam **associação, não causalidade**;",
        "Os fatores estão interligados (por exemplo, renda, nível e experiência). Um próximo passo é um **modelo "
        "multivariado** (regressão logística) para estimar o efeito de cada fator controlando os demais;",
        "A avaliação de desempenho tem só duas notas, o que limita a análise da Q5.",
    ])


def escrever_referencias(doc, e: Escritor, R: dict) -> None:
    titulo = _achar(doc, "6. Referências e Apêndices")
    titulo.paragraph_format.page_break_before = True
    e.posicionar(titulo)
    e.titulo2("6.1 Referências")
    refs = [
        f"IBM. **IBM HR Analytics Employee Attrition & Performance**. Kaggle, 2017. Disponível em: {config.URL_DATASET}. "
        "Acesso em: out. 2026.",
        "COHEN, J. **Statistical power analysis for the behavioral sciences**. 2. ed. Hillsdale: Lawrence Erlbaum, 1988.",
        "HOLM, S. A simple sequentially rejective multiple test procedure. **Scandinavian Journal of Statistics**, v. 6, "
        "n. 2, p. 65-70, 1979.",
        "NEWCOMBE, R. G. Interval estimation for the difference between independent proportions: comparison of eleven "
        "methods. **Statistics in Medicine**, v. 17, n. 8, p. 873-890, 1998.",
        "WELCH, B. L. The generalization of \"Student's\" problem when several different population variances are "
        "involved. **Biometrika**, v. 34, n. 1-2, p. 28-35, 1947.",
        "WILSON, E. B. Probable inference, the law of succession, and statistical inference. **Journal of the American "
        "Statistical Association**, v. 22, n. 158, p. 209-212, 1927.",
        "McKINNEY, W. Data structures for statistical computing in Python. In: **Proceedings of the 9th Python in "
        "Science Conference**, 2010. p. 56-61.",
        "VIRTANEN, P. et al. SciPy 1.0: fundamental algorithms for scientific computing in Python. **Nature Methods**, "
        "v. 17, p. 261-272, 2020.",
        "SEABOLD, S.; PERKTOLD, J. statsmodels: econometric and statistical modeling with Python. In: **Proceedings of "
        "the 9th Python in Science Conference**, 2010.",
        "HUNTER, J. D. Matplotlib: a 2D graphics environment. **Computing in Science & Engineering**, v. 9, n. 3, "
        "p. 90-95, 2007.",
        "WASKOM, M. L. seaborn: statistical data visualization. **Journal of Open Source Software**, v. 6, n. 60, "
        "p. 3021, 2021.",
    ]
    for ref in refs:
        e.paragrafo(ref, recuo=False, espaco=1.0, depois=8, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)

    e.titulo2("6.2 Código-fonte e reprodutibilidade")
    e.paragrafo(
        f"Repositório do projeto: **{URL_REPOSITORIO}**. Para reproduzir a análise: criar o ambiente virtual, instalar "
        f"*requirements.txt* e executar *python run_all.py*. Os notebooks *01_contexto_e_limpeza*, "
        f"*02_analise_descritiva* e *03_analise_inferencial* documentam cada decisão.", recuo=False)

    e.titulo2("Apêndice A – Verificação dos pressupostos")
    pres = pd.DataFrame([
        ["Todos", "Independência das observações", "1 linha = 1 funcionário; 0 duplicatas", "Atendido"],
        ["Qui-Quadrado (Q1, Q3, Q4)", "Frequências esperadas ≥ 5", "Menor esperada = 10,2", "Atendido"],
        ["Teste t (Q2)", "Normalidade (Shapiro-Wilk)", "Rejeitada (assimetria 1,9 e 1,3)", "Contornado pelo TCL (n = 237 e 1.233) e pela escala log"],
        ["Teste t (Q2)", "Igualdade de variâncias (Levene)", "Rejeitada (p < 0,001)", "Usado t de Welch"],
        ["Teste t (Q5a)", "Igualdade de variâncias (Levene)", "Rejeitada (p < 0,001)", "Usado t de Welch"],
        ["ANOVA (Q5b)", "Igualdade de variâncias (Levene)", "p = 0,72", "Atendido: ANOVA clássica"],
        ["ANOVA (Q5b)", "Normalidade dos resíduos", "Rejeitada", "Contornado pelo TCL (n ≥ 83 por grupo)"],
        ["Correlação (Q5c)", "Linearidade / outliers", "Renda assimétrica, sem padrão linear", "Usado Spearman"],
    ], columns=["Teste", "Pressuposto", "Verificação", "Resultado / tratamento"])
    e.tabela(pres, "Verificação dos pressupostos dos testes", larguras=[3.4, 3.8, 4.2, 4.6], tamanho=8)
    e.figura("fig12_pressupostos_qq", "Gráficos Q-Q da renda mensal (escala original e logarítmica)", largura=13,
             nota="Pontos próximos da reta indicam distribuição aproximadamente normal.")

    e.titulo2("Apêndice B – Dicionário de dados resumido")
    tip = R["tipagem"][["variavel", "descricao", "tipo_estatistico"]].rename(
        columns={"variavel": "Variável (original)", "descricao": "Descrição", "tipo_estatistico": "Tipo"})
    e.tabela(tip, "Variáveis analisadas, descrição e tipo estatístico", larguras=[5.0, 6.0, 5.0], tamanho=8)


# ---------------------------------------------------------------------------
# Sumário: preenchido com as páginas reais do PDF
# ---------------------------------------------------------------------------
def _titulos(doc: Document) -> list:
    """Lista (nível, texto) dos títulos que entram no sumário."""
    saida = []
    for par in doc.paragraphs:
        t = par.text.strip()
        if not t:
            continue
        if par.style.name == "Heading 1":
            saida.append((1, t))
        elif par.style.name == "Heading 2":
            saida.append((2, t))
    return saida


def _paginas_dos_titulos(pdf: Path, titulos: list) -> dict:
    """Procura cada título no texto do PDF (após o sumário) e devolve a página."""
    from pypdf import PdfReader
    paginas = [p.extract_text() or "" for p in PdfReader(str(pdf)).pages]
    normalizar = lambda s: re.sub(r"\s+", " ", s).strip()
    paginas = [normalizar(p) for p in paginas]
    inicio = next(i for i, p in enumerate(paginas) if "SUMÁRIO" in p) + 1
    resultado, atual = {}, inicio
    for _, texto in titulos:
        alvo = normalizar(texto)
        for i in range(atual, len(paginas)):
            if alvo in paginas[i]:
                resultado[texto] = i + 1
                atual = i
                break
    return resultado


def _reescrever_sumario(doc: Document, titulos: list, paginas: dict) -> None:
    """Substitui as entradas do sumário do modelo por entradas estáticas com as páginas certas."""
    sdt = next(doc.element.body.iter(qn("w:sdtContent")))
    modelo = sdt.find(qn("w:p"))
    for p in list(sdt):
        sdt.remove(p)
    for nivel, texto in titulos:
        p = copy.deepcopy(modelo)
        for filho in list(p):
            if filho.tag != qn("w:pPr"):
                p.remove(filho)
        ppr = p.find(qn("w:pPr"))
        espaco = ppr.find(qn("w:spacing"))
        espaco.set(qn("w:after"), "0")
        espaco.set(qn("w:before"), "40" if nivel == 2 else "100")
        for tab in ppr.iter(qn("w:tab")):
            tab.set(qn("w:pos"), "9060")  # tabulação à direita no limite da margem
        ind = ppr.find(qn("w:ind"))
        if nivel == 2:
            if ind is None:
                ind = OxmlElement("w:ind"); ppr.append(ind)
            ind.set(qn("w:left"), "360")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        if nivel == 1:
            rpr.append(OxmlElement("w:b"))
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "24"); rpr.append(sz)
        r.append(rpr)
        t1 = OxmlElement("w:t"); t1.text = texto; t1.set(qn("xml:space"), "preserve")
        r.append(t1)
        r.append(OxmlElement("w:tab"))
        t2 = OxmlElement("w:t"); t2.text = str(paginas.get(texto, ""))
        r.append(t2)
        p.append(r)
        sdt.append(p)


# ---------------------------------------------------------------------------
# Conversão para PDF
# ---------------------------------------------------------------------------
def converter_pdf(docx: Path, pasta_saida: Path) -> Path:
    """Converte DOCX → PDF com o LibreOffice em modo headless."""
    exe = shutil.which("libreoffice") or shutil.which("soffice")
    if exe is None:
        raise RuntimeError("LibreOffice não encontrado: instale-o para gerar o PDF.")
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(pasta_saida), str(docx)],
                   check=True, capture_output=True, timeout=300)
    return pasta_saida / (docx.stem + ".pdf")


def gerar_relatorio() -> None:
    R = carregar_resultados()
    doc = Document(str(config.MODELO_RELATORIO))
    _corrigir_margens(doc)
    _ajustar_capa(doc)
    _justificar_folha_de_rosto(doc)

    # remove os subtítulos vazios do modelo ("1.1.", "2.1." ...); o conteúdo cria os seus
    for par in list(doc.paragraphs):
        if par.style.name == "Heading 2" and re.fullmatch(r"\d\.\d\.?\s*", par.text.strip() + " "):
            par._p.getparent().remove(par._p)

    e = Escritor(doc)
    escrever_resumo(doc, e, R)
    escrever_introducao(doc, e, R)
    escrever_metodologia(doc, e, R)
    escrever_descritiva(doc, e, R)
    escrever_inferencial(doc, e, R)
    escrever_recomendacoes(doc, e, R)
    escrever_referencias(doc, e, R)

    # Sumário: gera o PDF, lê em que página cada título caiu e reescreve o sumário.
    # Repete até as páginas estabilizarem (o próprio sumário pode mudar de tamanho).
    titulos = _titulos(doc)
    paginas = {}
    for _ in range(4):
        _reescrever_sumario(doc, titulos, paginas)
        doc.save(SAIDA_DOCX)
        pdf = converter_pdf(SAIDA_DOCX, config.PASTA_REPORTS)
        novas = _paginas_dos_titulos(pdf, titulos)
        if novas == paginas:
            break
        paginas = novas
    faltando = [t for _, t in titulos if t not in paginas]
    if faltando:
        print("Aviso: títulos não localizados no PDF:", faltando)
    print(f"Relatório gerado: {SAIDA_DOCX.relative_to(config.RAIZ)} e {SAIDA_PDF.relative_to(config.RAIZ)}")
    print(f"{e.n_fig} figuras e {e.n_tab} tabelas")


if __name__ == "__main__":
    gerar_relatorio()
